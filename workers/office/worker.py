from __future__ import annotations

import csv
import io
import json
import os
import stat
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import Workbook, load_workbook


class WorkerError(ValueError):
    """A request cannot be handled safely or deterministically."""


def _absolute_dir(value: Any, label: str) -> Path:
    if not isinstance(value, str) or not value or not Path(value).is_absolute():
        raise WorkerError(f"{label} must be an absolute path")
    path = Path(value)
    _reject_link_components(path)
    if not path.is_dir():
        raise WorkerError(f"{label} must be an existing directory")
    return path.resolve(strict=True)


def _reject_link_components(path: Path) -> None:
    """Reject symlinks and Windows junctions, including dangling links."""
    absolute = Path(os.path.abspath(path))
    current = Path(absolute.anchor)
    for part in absolute.parts[1:]:
        current = current / part
        try:
            info = current.lstat()
        except FileNotFoundError:
            continue
        if stat.S_ISLNK(info.st_mode) or getattr(info, "st_file_attributes", 0) & getattr(stat, "FILE_ATTRIBUTE_REPARSE_POINT", 0x400):
            raise WorkerError(f"symlink or junction is not allowed: {current}")


def _relative_path(root: Path, value: Any, label: str) -> Path:
    if not isinstance(value, str) or not value:
        raise WorkerError(f"{label} must be a non-empty relative path")
    candidate = Path(value)
    if candidate.is_absolute() or any(part == ".." for part in candidate.parts):
        raise WorkerError(f"{label} must stay within workspace")
    path = root / candidate
    _reject_link_components(path)
    try:
        resolved = path.resolve(strict=True)
        resolved.relative_to(root)
    except (FileNotFoundError, ValueError) as exc:
        raise WorkerError(f"{label} does not exist within workspace") from exc
    _reject_link_components(path)
    return resolved


def _prepare_output(root: Path, value: Any) -> Path:
    if not isinstance(value, str) or not value or not Path(value).is_absolute():
        raise WorkerError("outputDir must be an absolute path")
    raw = Path(os.path.abspath(value))
    if any(part == ".." for part in Path(value).parts):
        raise WorkerError("outputDir must stay within workspace")
    try:
        raw.relative_to(root)
    except ValueError as exc:
        raise WorkerError("outputDir must be inside workspace") from exc
    _reject_link_components(raw)
    raw.mkdir(parents=True, exist_ok=True)
    _reject_link_components(raw)
    resolved = raw.resolve(strict=True)
    try:
        resolved.relative_to(root)
    except ValueError as exc:
        raise WorkerError("outputDir must be inside workspace") from exc
    return resolved


def _write_exclusive(path: Path, data: bytes) -> None:
    flags = os.O_WRONLY | os.O_CREAT | os.O_EXCL
    fd = os.open(path, flags, 0o600)
    try:
        with os.fdopen(fd, "wb") as stream:
            stream.write(data)
    except Exception:
        try:
            path.unlink()
        except OSError:
            pass
        raise


def _artifacts(output: Path, files: list[tuple[str, bytes, str, str | None]]) -> list[dict[str, Any]]:
    destinations = [output / filename for filename, _, _, _ in files]
    for dest in destinations:
        if dest.exists() or dest.is_symlink():
            raise WorkerError(f"output already exists: {dest}")
    results = []
    for (filename, data, mime, preview), dest in zip(files, destinations, strict=True):
        _write_exclusive(dest, data)
        item: dict[str, Any] = {"path": str(dest), "name": filename, "mime": mime}
        if preview is not None:
            item["preview"] = preview
        results.append(item)
    return results


def _text(value: Any, label: str, *, default: str | None = None) -> str:
    if value is None and default is not None:
        return default
    if not isinstance(value, str):
        raise WorkerError(f"{label} must be a string")
    return value


def _inventory(root: Path, output: Path, inputs: list[Any], options: dict[str, Any]) -> dict[str, Any]:
    directory_value = options.get("directory", inputs[0] if inputs else ".")
    directory = _relative_path(root, directory_value, "inventory directory")
    if not directory.is_dir():
        raise WorkerError("inventory directory must be a directory")
    records: list[tuple[str, int]] = []
    for current, dirs, files in os.walk(directory, followlinks=False):
        base = Path(current)
        for name in list(dirs):
            child = base / name
            _reject_link_components(child)
        for name in files:
            child = base / name
            _reject_link_components(child)
            if not child.is_file():
                raise WorkerError(f"inventory encountered a non-regular file: {child}")
            if child == output or output in child.parents:
                continue
            records.append((child.relative_to(root).as_posix(), child.stat().st_size))
    records.sort(key=lambda row: row[0].casefold())
    table = io.StringIO(newline="")
    writer = csv.writer(table, lineterminator="\n")
    writer.writerow(["path", "size_bytes"])
    writer.writerows(records)
    lines = ["# 文件清单", "", f"文件数：{len(records)}", "", "| 路径 | 大小（字节） |", "| --- | ---: |"]
    lines.extend(f"| `{path.replace('|', '\\|')}` | {size} |" for path, size in records)
    md = "\n".join(lines) + "\n"
    artifacts = _artifacts(output, [
        ("inventory.csv", table.getvalue().encode("utf-8-sig"), "text/csv", None),
        ("inventory.md", md.encode("utf-8"), "text/markdown", md[:2000]),
    ])
    return {"ok": True, "summary": f"已清点 {len(records)} 个文件", "artifacts": artifacts, "details": {"fileCount": len(records)}}


def _blank_record(values: list[Any] | tuple[Any, ...]) -> bool:
    return not values or all(value is None or (isinstance(value, str) and not value.strip()) for value in values)


def _check_formula_cells(path: Path, row_number: int, cells: list[Any] | tuple[Any, ...]) -> None:
    for column_number, cell in enumerate(cells, start=1):
        if cell.data_type == "f":
            raise WorkerError(
                f"{path.name} 第 {row_number} 行第 {column_number} 列包含公式；请导出为纯值后核对"
            )


def _load_table(path: Path) -> tuple[list[str], list[tuple[int, dict[str, Any]]]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            reader = csv.reader(stream)
            try:
                headers = next(reader)
            except StopIteration as exc:
                raise WorkerError(f"input table is empty: {path.name}") from exc
            rows: list[tuple[int, dict[str, Any]]] = []
            for values in reader:
                row_number = reader.line_num
                if _blank_record(values):
                    continue
                if len(values) != len(headers):
                    raise WorkerError(
                        f"{path.name} 第 {row_number} 行列数为 {len(values)}，表头列数为 {len(headers)}"
                    )
                rows.append((row_number, dict(zip(headers, values, strict=True))))
    elif suffix == ".xlsx":
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            sheet = book.active
            iterator = sheet.iter_rows()
            try:
                header_cells = next(iterator)
            except StopIteration as exc:
                raise WorkerError(f"input table is empty: {path.name}") from exc
            _check_formula_cells(path, 1, header_cells)
            headers = [str(cell.value) if cell.value is not None else "" for cell in header_cells]
            rows = []
            for row_number, cells in enumerate(iterator, start=2):
                _check_formula_cells(path, row_number, cells)
                values = [cell.value for cell in cells]
                if _blank_record(values):
                    continue
                if len(values) != len(headers):
                    raise WorkerError(
                        f"{path.name} 第 {row_number} 行列数为 {len(values)}，表头列数为 {len(headers)}"
                    )
                rows.append((row_number, dict(zip(headers, values, strict=True))))
        finally:
            book.close()
    else:
        raise WorkerError(f"unsupported table format: {path.suffix}; expected CSV or XLSX")
    if not headers or any(not header for header in headers):
        raise WorkerError(f"table headers must be non-empty: {path.name}")
    if len(headers) != len(set(headers)):
        raise WorkerError(f"table headers must be unique: {path.name}")
    return headers, rows


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def _reconcile(root: Path, output: Path, inputs: list[Any], options: dict[str, Any]) -> dict[str, Any]:
    if len(inputs) != 2:
        raise WorkerError("reconcile requires exactly two input files: baseline then comparison")
    left_path = _relative_path(root, inputs[0], "inputs[0]")
    right_path = _relative_path(root, inputs[1], "inputs[1]")
    if not left_path.is_file() or not right_path.is_file():
        raise WorkerError("reconcile inputs must be files")
    left_headers, left_rows = _load_table(left_path)
    right_headers, right_rows = _load_table(right_path)
    key_option = options.get("key")
    if key_option is None:
        key = left_headers[0]
    elif isinstance(key_option, int) and not isinstance(key_option, bool):
        if key_option < 0 or key_option >= len(left_headers):
            raise WorkerError("options.key column index is out of range")
        key = left_headers[key_option]
    else:
        key = _text(key_option, "options.key")
    if key not in left_headers or key not in right_headers:
        raise WorkerError(f"key column must exist in both tables: {key}")
    columns = list(dict.fromkeys([*left_headers, *right_headers]))
    left_groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    right_groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for side_path, rows, groups in [
        (left_path, left_rows, left_groups),
        (right_path, right_rows, right_groups),
    ]:
        for row_number, row in rows:
            entry_key = _cell(row.get(key))
            if not entry_key.strip():
                raise WorkerError(f"{side_path.name} 第 {row_number} 行主键 `{key}` 不能为空")
            groups[entry_key].append(row)
    left_dupes = {k: len(v) for k, v in left_groups.items() if len(v) > 1}
    right_dupes = {k: len(v) for k, v in right_groups.items() if len(v) > 1}
    differences: list[tuple[str, str, str, str]] = []
    keys = sorted(set(left_groups) | set(right_groups))
    for entry_key in keys:
        lrows, rrows = left_groups.get(entry_key, []), right_groups.get(entry_key, [])
        if len(lrows) > 1 or len(rrows) > 1:
            if len(lrows) > 1:
                differences.append(("duplicate_baseline", entry_key, json.dumps(lrows, ensure_ascii=False, default=_cell), ""))
            if len(rrows) > 1:
                differences.append(("duplicate_comparison", entry_key, "", json.dumps(rrows, ensure_ascii=False, default=_cell)))
            continue
        if not lrows:
            differences.append(("added", entry_key, "", json.dumps(rrows[0], ensure_ascii=False, default=_cell)))
        elif not rrows:
            differences.append(("missing", entry_key, json.dumps(lrows[0], ensure_ascii=False, default=_cell), ""))
        elif any(_cell(lrows[0].get(column)) != _cell(rrows[0].get(column)) for column in columns):
            differences.append(("modified", entry_key, json.dumps(lrows[0], ensure_ascii=False, default=_cell), json.dumps(rrows[0], ensure_ascii=False, default=_cell)))

    csv_stream = io.StringIO(newline="")
    csv_writer = csv.writer(csv_stream, lineterminator="\n")
    def safe_csv_cell(value: str) -> tuple[str, bool]:
        if value.startswith(("=", "+", "-", "@", "\t", "\r")):
            return "'" + value, True
        return value, False

    raw_csv_rows = [["type", key, "baseline", "comparison"], *[list(row) for row in differences]]
    csv_rows: list[list[str]] = []
    protected_csv_cells = 0
    for row in raw_csv_rows:
        safe_row = []
        for value in row:
            safe, protected = safe_csv_cell(value)
            safe_row.append(safe)
            protected_csv_cells += int(protected)
        csv_rows.append(safe_row)
    csv_writer.writerows(csv_rows)
    csv_preview = csv_stream.getvalue()[:3000]
    xlsx_buffer = io.BytesIO()
    book = Workbook()
    sheet = book.active
    sheet.title = "differences"
    for row_number, values in enumerate([["type", key, "baseline", "comparison"], *[list(row) for row in differences]], start=1):
        for column_number, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_number, column=column_number, value=str(value))
            cell.data_type = "s"
    book.save(xlsx_buffer)
    book.close()

    counts = Counter(row[0] for row in differences)
    lines = ["# 核对报告", "", f"基准文件：`{inputs[0]}`  ", f"对比文件：`{inputs[1]}`  ", f"核对键：`{key}`", "", "## 结果", ""]
    for kind, label in [("added", "新增"), ("missing", "缺失"), ("modified", "修改"), ("duplicate_baseline", "基准表重复键"), ("duplicate_comparison", "对比表重复键")]:
        lines.append(f"- {label}：{counts[kind]}")
    for title, dupes in [("基准表重复键详情", left_dupes), ("对比表重复键详情", right_dupes)]:
        lines.extend(["", f"## {title}", ""])
        lines.extend([f"- `{entry_key}`：{count} 条" for entry_key, count in sorted(dupes.items())] or ["- 无"])
    md = "\n".join(lines) + "\n"
    artifacts = _artifacts(output, [
        ("differences.xlsx", xlsx_buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", None),
        ("differences.csv", csv_stream.getvalue().encode("utf-8-sig"), "text/csv", csv_preview),
        ("report.md", md.encode("utf-8"), "text/markdown", md[:2000]),
    ])
    details = {
        "added": counts["added"], "missing": counts["missing"], "modified": counts["modified"],
        "duplicateBaseline": left_dupes, "duplicateComparison": right_dupes,
        "csvFormulaProtection": {
            "strategy": "prefix apostrophe when a cell starts with =, +, -, @, tab, or carriage return",
            "protectedCells": protected_csv_cells,
        },
    }
    return {"ok": True, "summary": f"核对完成：新增 {counts['added']}、缺失 {counts['missing']}、修改 {counts['modified']}，重复键 {counts['duplicate_baseline'] + counts['duplicate_comparison']} 组", "artifacts": artifacts, "details": details}


def _report(output: Path, options: dict[str, Any]) -> dict[str, Any]:
    title = _text(options.get("title", ""), "options.title")
    content = _text(options.get("content", ""), "options.content")
    md = (f"# {title}\n\n" if title else "") + content
    if not md.endswith("\n"):
        md += "\n"
    document = Document()
    if title:
        document.add_heading(title, level=1)
    for paragraph in content.splitlines() or ([""] if content else []):
        document.add_paragraph(paragraph)
    docx_buffer = io.BytesIO()
    document.save(docx_buffer)
    artifacts = _artifacts(output, [
        ("report.md", md.encode("utf-8"), "text/markdown", md[:2000]),
        ("report.docx", docx_buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", None),
    ])
    return {"ok": True, "summary": "报告已生成", "artifacts": artifacts, "details": {"title": title, "contentCharacters": len(content)}}


def run_request(request: dict[str, Any]) -> dict[str, Any]:
    operation = request.get("operation")
    if operation not in {"inventory", "reconcile", "report"}:
        raise WorkerError("operation must be inventory, reconcile, or report")
    workspace = _absolute_dir(request.get("workspace"), "workspace")
    output = _prepare_output(workspace, request.get("outputDir"))
    inputs = request.get("inputs", [])
    options = request.get("options", {})
    if not isinstance(inputs, list) or not isinstance(options, dict):
        raise WorkerError("inputs must be an array and options must be an object")
    if operation == "inventory":
        return _inventory(workspace, output, inputs, options)
    if operation == "reconcile":
        return _reconcile(workspace, output, inputs, options)
    return _report(output, options)
