from __future__ import annotations

import csv
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from workers.office.worker import WorkerError, run_request


class OfficeWorkerTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name).resolve()
        (self.root / "base.csv").write_text("id,name\n1,A\n2,B\n", encoding="utf-8")
        (self.root / "current.csv").write_text("id,name\n1,A\n2,C\n3,D\n", encoding="utf-8")

    def tearDown(self) -> None:
        self.temp.cleanup()

    def request(self, operation: str, **kwargs: object) -> dict[str, object]:
        return run_request({
            "operation": operation,
            "workspace": str(self.root),
            "outputDir": str(self.root / "out"),
            "inputs": [],
            "options": {},
            **kwargs,
        })

    def test_reconcile_detects_add_missing_and_modified_rows(self) -> None:
        result = self.request("reconcile", inputs=["base.csv", "current.csv"])
        self.assertEqual(result["details"]["added"], 1)
        self.assertEqual(result["details"]["missing"], 0)
        self.assertEqual(result["details"]["modified"], 1)
        with (self.root / "out" / "differences.csv").open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.DictReader(stream))
        self.assertEqual([row["type"] for row in rows], ["modified", "added"])
        self.assertEqual({artifact["name"] for artifact in result["artifacts"]}, {"differences.xlsx", "differences.csv", "report.md"})

    def test_duplicate_keys_are_reported_and_not_overwritten(self) -> None:
        (self.root / "base.csv").write_text("id,name\n1,A\n1,B\n", encoding="utf-8")
        result = self.request("reconcile", inputs=["base.csv", "current.csv"])
        self.assertEqual(result["details"]["duplicateBaseline"], {"1": 2})
        self.assertEqual(result["details"]["modified"], 0)
        md = (self.root / "out" / "report.md").read_text(encoding="utf-8")
        self.assertIn("基准表重复键详情", md)
        self.assertIn("`1`：2 条", md)

    def test_rejects_csv_row_with_extra_column_and_reports_line_number(self) -> None:
        (self.root / "base.csv").write_text("id,name\n1,A\n2,B,extra\n", encoding="utf-8")
        with self.assertRaisesRegex(WorkerError, r"base\.csv 第 3 行列数为 3，表头列数为 2"):
            self.request("reconcile", inputs=["base.csv", "current.csv"])

    def test_ignores_blank_csv_records(self) -> None:
        (self.root / "base.csv").write_text("id,name\n1,A\n,  \n2,B\n", encoding="utf-8")
        result = self.request("reconcile", inputs=["base.csv", "current.csv"])
        self.assertEqual(result["details"]["modified"], 1)

    def test_rejects_xlsx_formula_and_asks_for_values_export(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["id", "name"])
        sheet.append(["=1+1", "A"])
        workbook.save(self.root / "formula.xlsx")
        workbook.close()
        with self.assertRaisesRegex(WorkerError, r"formula\.xlsx 第 2 行第 1 列包含公式；请导出为纯值后核对"):
            self.request("reconcile", inputs=["formula.xlsx", "current.csv"])

    def test_rejects_empty_primary_key_with_row_number(self) -> None:
        (self.root / "base.csv").write_text("id,name\n,A\n1,ok\n", encoding="utf-8")
        with self.assertRaisesRegex(WorkerError, r"base\.csv 第 2 行主键 `id` 不能为空"):
            self.request("reconcile", inputs=["base.csv", "current.csv"])

    def test_formula_style_key_is_written_as_text_to_xlsx_and_protected_in_csv(self) -> None:
        (self.root / "base.csv").write_text("id,name\n=1+1,before\n", encoding="utf-8")
        (self.root / "current.csv").write_text("id,name\n=1+1,after\n", encoding="utf-8")
        result = self.request("reconcile", inputs=["base.csv", "current.csv"])

        self.assertEqual(result["details"]["csvFormulaProtection"]["protectedCells"], 1)
        self.assertIn("apostrophe", result["details"]["csvFormulaProtection"]["strategy"])
        with (self.root / "out" / "differences.csv").open(encoding="utf-8-sig", newline="") as stream:
            difference = next(csv.DictReader(stream))
        self.assertEqual(difference["id"], "'=1+1")

        workbook = load_workbook(self.root / "out" / "differences.xlsx", data_only=False)
        try:
            sheet = workbook.active
            cells = list(sheet.iter_rows())
            self.assertTrue(cells)
            self.assertTrue(all(cell.data_type == "s" for row in cells for cell in row))
            self.assertEqual(cells[1][1].value, "=1+1")
            self.assertIn('"id": "=1+1"', cells[1][2].value)
        finally:
            workbook.close()

    def test_rejects_parent_traversal(self) -> None:
        with self.assertRaisesRegex(WorkerError, "workspace"):
            self.request("reconcile", inputs=["../outside.csv", "current.csv"])

    def test_output_dir_must_be_inside_workspace(self) -> None:
        with self.assertRaisesRegex(WorkerError, "inside workspace"):
            run_request({"operation": "report", "workspace": str(self.root), "outputDir": str(self.root.parent / "outside"), "options": {}})

    def test_existing_output_is_never_overwritten(self) -> None:
        out = self.root / "out"
        out.mkdir()
        existing = out / "differences.xlsx"
        existing.write_bytes(b"keep")
        with self.assertRaisesRegex(WorkerError, "already exists"):
            self.request("reconcile", inputs=["base.csv", "current.csv"])
        self.assertEqual(existing.read_bytes(), b"keep")


if __name__ == "__main__":
    unittest.main()
